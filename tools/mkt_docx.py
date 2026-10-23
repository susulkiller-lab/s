"""마케팅 노출 리포트 DOCX 출력(mkt_report.py가 호출). 양식: 맑은 고딕, A4, 여백 2cm, 표 머리 음영.

필요: python-docx (`py -m pip install python-docx`). 없으면 mkt_report.py는 DOCX만 건너뛴다.
"""
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Twips

FONT = "맑은 고딕"
NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x59, 0x59, 0x59)
HEAD_FILL, US_FILL, BOX_FILL, LINE = "DCE3EE", "FFF2CC", "F2F5FA", "BFBFBF"
BODY = 8.5
WIDTH = 9638  # A4 폭 - 좌우 여백(트윕)
ENGINE = {"chatgpt": "ChatGPT", "gemini": "Gemini", "perplexity": "Perplexity", "claude": "Claude"}


def dot(s):
    """2026-10-23 → 2026. 10. 23."""
    return re.sub(r"(\d{4})-(\d{2})-(\d{2})", lambda m: f"{m.group(1)}. {int(m.group(2))}. {int(m.group(3))}.", s)


def run(par, text, size=BODY, bold=False, color=None):
    r = par.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    if color is not None:
        r.font.color.rgb = color
    return r


def para(doc, text="", size=BODY, bold=False, color=None, align=None, before=0, after=60, style=None):
    p = doc.add_paragraph(style=style)
    if text:
        run(p, text, size, bold, color)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before / 20), Pt(after / 20)
    if align:
        p.alignment = align
    return p


def heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    run(p, text, 12, True, NAVY)
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Twips(320), Twips(120)
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:color", "1F3864"), ("w:sz", "8"), ("w:space", "2")):
        b.set(qn(k), v)
    bdr.append(b)
    p._p.get_or_add_pPr().append(bdr)


def note(doc, text):
    return para(doc, text, color=GRAY, after=80)


def subhead(doc, text):
    return para(doc, text, 10, True, before=160, after=80)


def bullet(doc, text, level=0):
    return para(doc, text, style="List Bullet 2" if level else "List Bullet", after=40)


def _cell_props(cell, width, fill=None, margin=(50, 80)):
    tcPr = cell._tc.get_or_add_tcPr()
    w = OxmlElement("w:tcW")
    w.set(qn("w:type"), "dxa")
    w.set(qn("w:w"), str(width))
    tcPr.append(w)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{side}")
        for k, v in (("w:val", "single"), ("w:color", LINE), ("w:sz", "4")):
            e.set(qn(k), v)
        borders.append(e)
    tcPr.append(borders)
    if fill:
        shd = OxmlElement("w:shd")
        for k, v in (("w:fill", fill), ("w:color", "auto"), ("w:val", "clear")):
            shd.set(qn(k), v)
        tcPr.append(shd)
    mar = OxmlElement("w:tcMar")
    for side, v in (("top", margin[0]), ("left", margin[1]), ("bottom", margin[0]), ("right", margin[1])):
        e = OxmlElement(f"w:{side}")
        e.set(qn("w:type"), "dxa")
        e.set(qn("w:w"), str(v))
        mar.append(e)
    tcPr.append(mar)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for p in cell.paragraphs:
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0


def _table(doc, ncols, widths):
    t = doc.add_table(rows=0, cols=ncols)
    t.autofit = False
    tblPr = t._tbl.tblPr
    tw = OxmlElement("w:tblW")
    tw.set(qn("w:type"), "dxa")
    tw.set(qn("w:w"), str(sum(widths)))
    tblPr.append(tw)
    grid = t._tbl.tblGrid
    for i, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(widths[i]))
    return t


def table(doc, headers, rows, widths, align=None, highlight=()):
    """headers: 열 제목, rows: 문자열 목록의 목록, widths: 트윕(합계 WIDTH), align: 열별 'l'/'c', highlight: 강조 행 번호."""
    align = align or ["c"] * len(headers)
    t = _table(doc, len(headers), widths)
    hdr = t.add_row()
    tr = hdr._tr.get_or_add_trPr()
    tr.append(OxmlElement("w:tblHeader"))
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        _cell_props(c, widths[i], HEAD_FILL)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(c.paragraphs[0], h, bold=True)
    for ri, row in enumerate(rows):
        r = t.add_row()
        hl = ri in highlight
        for i, val in enumerate(row):
            c = r.cells[i]
            _cell_props(c, widths[i], US_FILL if hl else None)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if align[i] == "l" else WD_ALIGN_PARAGRAPH.CENTER
            run(c.paragraphs[0], str(val), bold=hl)
    para(doc, after=0)
    return t


def boxes(doc, items):
    """결론 요약 칸: (제목, 큰 숫자, 설명)."""
    w = [WIDTH // len(items)] * len(items)
    w[0] += WIDTH - sum(w)
    t = _table(doc, len(items), w)
    r = t.add_row()
    for i, (label, big, sub) in enumerate(items):
        c = r.cells[i]
        _cell_props(c, w[i], BOX_FILL, (120, 120))
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(c.paragraphs[0], label, color=GRAY)
        p2 = c.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = p2.paragraph_format.space_after = Twips(40)
        run(p2, big, 16, True, NAVY)
        p3 = c.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run(p3, sub, color=GRAY)
    para(doc, after=0)


def footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, "- ", 8, color=RGBColor(0x80, 0x80, 0x80))
    r = run(p, "", 8, color=RGBColor(0x80, 0x80, 0x80))
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            f = OxmlElement("w:fldChar")
            f.set(qn("w:fldCharType"), kind)
            r._r.append(f)
        else:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
    run(p, " -", 8, color=RGBColor(0x80, 0x80, 0x80))


def build(R, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Twips(1134))
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(BODY)
    st.paragraph_format.space_before = st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.15
    footer(sec)

    para(doc, "진앤솔법률사무소 마케팅 노출 리포트", 17, True, NAVY, WD_ALIGN_PARAGRAPH.CENTER, after=40)
    para(doc, f"측정일 {dot(R['date'])}", 10, color=GRAY, align=WD_ALIGN_PARAGRAPH.CENTER, after=200)
    meta = [("비교 기준", dot(R["prev"]) + " 측정" if R["prev"] else "직전 측정 없음(이번이 기준선)"),
            ("측정 조건", R["conditions"])]
    if R["missing_channels"]:
        meta.append(("미측정 채널", f"{', '.join(R['missing_channels'])} (종합지수는 측정된 채널 가중치로만 환산)"))
    table(doc, ["항목", "내용"], meta, [1800, WIDTH - 1800], ["c", "l"])

    # 1. 결론
    heading(doc, "1. 결론")
    S = R["summary"]
    items = [("종합 노출지수", str(S["index"]), f"비교군 {S['n']}곳 중 {S['pos']}위")]
    if S.get("share") is not None:
        items.append(("노출 점유율", f"{S['share']}%", "비교군 지수 합 대비"))
    if S.get("lead"):
        items.append(("1위 대비 격차" if S["lead_is_first"] else "2위와의 차이", str(S["gap"]),
                      f"{'1위' if S['lead_is_first'] else '2위'} {S['lead']} {S['lead_index']}"))
    boxes(doc, items)
    if R.get("channel_rank"):
        subhead(doc, "채널별 우리 위치")
        for n, (pos, cnt, mine, top, ties) in R["channel_rank"].items():
            rank = f"공동 {pos}위({ties}곳)" if ties > 1 else f"{pos}위"
            bullet(doc, f"{n}: {cnt}곳 중 {rank}, 우리 {mine} / 1위 {top}")

    # 2. 채널별 점수
    heading(doc, "2. 채널별 점수")
    note(doc, "점수 범위 0~100. 진앤솔 행은 음영 표시." + (" 괄호는 직전 대비." if R["prev"] else ""))
    names = R["channel_names"]
    w0 = 1838
    table(doc, ["업체"] + names + ["종합"],
          [[k + ("*" if us else "")] + vals + [idx] for k, us, vals, idx in R["channels"]],
          [w0] + [(WIDTH - w0) // (len(names) + 1)] * len(names) + [WIDTH - w0 - (WIDTH - w0) // (len(names) + 1) * len(names)],
          ["l"] + ["c"] * (len(names) + 1), {i for i, r in enumerate(R["channels"]) if r[1]})

    subhead(doc, "업체별 노출 키워드 수")
    table(doc, ["업체", "파워링크 노출 키워드 수", "구글 지도 결과 노출 키워드 수"],
          [[k, "-" if a is None else a, "-" if g is None else g] for k, _, a, g in R["counts"]],
          [3212, 3213, 3213], ["l", "c", "c"], {i for i, r in enumerate(R["counts"]) if r[1]})
    if R["advertisers"]:
        subhead(doc, "파워링크 광고주(노출 키워드 수)")
        table(doc, ["광고주 도메인", "노출 키워드 수"], [[d, n] for d, n in R["advertisers"]], [6425, 3213], ["l", "c"])

    if R["ai_table"]:
        subhead(doc, "AI 검색 엔진별 언급")
        en = R["ai_engines"]
        shown = [r for r in R["ai_table"] if r[1] or sum(r[2])]
        hidden = len(R["ai_table"]) - len(shown)
        note(doc, "엔진별 칸은 언급된 답변 수 / 전체 답변 수. 출처 인용은 답변 안에 그 업체 홈페이지·블로그 주소가 나온 건수."
             + (f" 언급이 없는 {hidden}곳은 생략." if hidden else ""))
        we = (WIDTH - 1838 - 1400 - 1400) // len(en)
        widths = [1838] + [we] * len(en) + [1400, WIDTH - 1838 - we * len(en) - 1400]
        table(doc, ["업체"] + [ENGINE.get(x, x) for x in en] + ["언급률", "출처 인용"],
              [[k] + [f"{c} / {R['ai_total'][x]}" for c, x in zip(cnt, en)] + [rate, cites] for k, _, cnt, rate, cites in shown],
              widths, ["l"] + ["c"] * (len(en) + 2), {i for i, r in enumerate(shown) if r[1]})
        if R["ai_us"]:
            subhead(doc, "진앤솔이 언급된 AI 답변")
            table(doc, ["엔진", "질문", "언급 순서"],
                  [[ENGINE.get(e, e), q, f"{pos}번째 / 언급 업체 {n}곳"] for e, q, pos, n in R["ai_us"]],
                  [1400, 5838, 2400], ["c", "l", "c"])

    # 3. 약점 목록
    heading(doc, "3. 약점 목록")
    note(doc, '경쟁사가 우리보다 위에 있는 키워드, 가중치 × 점수 차 순. 지표는 "1위 업체 / 우리" 순서로 표기.')
    if R["gaps"]:
        table(doc, ["채널", "키워드", "우리 순위", "1위 업체", "방문자리뷰", "사진", "톡톡", "쿠폰"], R["gaps"],
              [900, 2038, 900, 1300, 1300, 1000, 1100, 1100])
    else:
        bullet(doc, "해당 없음")

    # 4. 순위 변동
    heading(doc, "4. 우리 순위 변동")
    note(doc, "직전 대비, 3계단 이상 또는 노출·미노출 전환")
    for m in R["moves"]:
        bullet(doc, m)

    # 5. 운영 점검
    heading(doc, "5. 운영 점검")
    note(doc, "우리 측 감점 항목, 관측된 것만")
    for fl in R["flags"] or ["관측된 감점 항목 없음"]:
        bullet(doc, dot(fl))
        if fl.startswith("발행 글 누락"):
            for d, t in R["lost"]:
                bullet(doc, f'{dot(d)} 발행 "{t}"', 1)
    if R["place_ranks"]:
        subhead(doc, "우리 플레이스 순위(전체 업체 수 대비)")
        pr = R["place_ranks"]
        half = (len(pr) + 1) // 2
        rows = [list(pr[i]) + (list(pr[i + half]) if i + half < len(pr) else ["", "", ""]) for i in range(half)]
        table(doc, ["키워드", "우리 순위", "전체"] * 2, rows, [2419, 1200, 1200, 2419, 1200, 1200],
              ["l", "c", "c", "l", "c", "c"])

    # 6. 상관
    heading(doc, "6. 플레이스 순위와 공개지표의 상관")
    note(doc, "스피어만 상관계수. -1에 가까울수록 지표가 클수록 상위.")
    table(doc, ["지표", "상관계수", "표본"], [list(x) for x in R.get("corr", [])], [3212, 3213, 3213])
    note(doc, "상관관계이며 인과가 아닙니다. 같은 업체가 여러 키워드에 중복 집계되고, 4주 이상 누적해야 경향으로 읽을 수 있습니다.")

    # 7. 비교군 후보
    heading(doc, "7. 비교군 후보")
    note(doc, "우리·등록 업체·네트워크펌 제외, 상위 노출 가중 점수 순")
    if R["candidates"]:
        table(doc, ["업체명", "플레이스 ID", "점수"], [[n, pid, s] for n, pid, s in R["candidates"]],
              [6038, 2000, 1600], ["l", "c", "c"])
    else:
        bullet(doc, "후보 없음")

    # 8. 상담
    heading(doc, "8. 상담·수임 실적")
    note(doc, "data/consults.csv 입력분")
    for c in R["consults"]:
        bullet(doc, dot(c))

    # 9. 과제
    heading(doc, "9. 다음 과제 후보")
    note(doc, "위 관측 결과에서 기계적으로 도출, 최종 선택은 사무장님")
    for head, items in R["todo"]:
        bullet(doc, head)
        if items:
            bullet(doc, ", ".join(items), 1)

    doc.save(path)
